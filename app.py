#!/usr/bin/env python3
"""
CERT-X Web  -  browser-based SSL certificate expiry checker
-----------------------------------------------------------
Upload an Excel/CSV of website names (first column, or a column named
website/url/certificate). The server reads each certificate, classifies
expiry with the CERT-X rules, shows a color-coded table, and lets you
download an Excel report.

Run locally:
    pip install flask openpyxl requests truststore
    python app.py
    # open http://127.0.0.1:8000
"""
import io
import os
import csv
import threading
from concurrent.futures import ThreadPoolExecutor

from flask import Flask, request, render_template, send_file, jsonify

import cert_core

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    HAVE_OPENPYXL = True
except Exception:
    HAVE_OPENPYXL = False

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB uploads

COLS = ["Website", "Expiry Date", "Days Left", "Active Status",
        "Filter", "Action Status", "Issuer", "Issued On", "Error"]


# ---------- input parsing ----------
def read_sites_from_bytes(filename, data):
    rows = []
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        if not HAVE_OPENPYXL:
            raise RuntimeError("openpyxl not installed on the server.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        col_idx = 0
        keys = ("website", "url", "certificate", "site", "domain",
                "certificate name", "website / certificate name")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if row is None:
                continue
            if i == 0:
                header = [str(c).strip().lower() if c is not None else "" for c in row]
                found = False
                for key in keys:
                    if key in header:
                        col_idx = header.index(key); found = True; break
                if not found:
                    val = row[col_idx] if col_idx < len(row) else None
                    h = cert_core.normalize_host(val) if val is not None else None
                    if h: rows.append(h)
                continue
            val = row[col_idx] if col_idx < len(row) else None
            if val is None:
                continue
            h = cert_core.normalize_host(val)
            if h: rows.append(h)
    else:
        text = data.decode("utf-8-sig", "replace").splitlines()
        reader = csv.reader(text)
        first = True; col_idx = 0
        for r in reader:
            if not r:
                continue
            if first:
                first = False
                low = [c.strip().lower() for c in r]; picked = False
                for key in ("website", "url", "certificate", "site", "domain"):
                    if key in low:
                        col_idx = low.index(key); picked = True; break
                if picked:
                    continue
            val = r[col_idx] if col_idx < len(r) else None
            h = cert_core.normalize_host(val) if val else None
            if h: rows.append(h)
    # de-dup, preserve order
    seen = set(); uniq = []
    for h in rows:
        if h not in seen:
            seen.add(h); uniq.append(h)
    return uniq


# ---------- excel report ----------
def build_xlsx(results):
    wb = Workbook(); ws = wb.active; ws.title = "Certificates"
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    level_fill = {"green": PatternFill("solid", fgColor="DCFCE7"),
                  "yellow": PatternFill("solid", fgColor="FEF3C7"),
                  "red": PatternFill("solid", fgColor="FEE2E2")}
    level_font = {"green": Font(color="166534"), "yellow": Font(color="92400E"),
                  "red": Font(color="991B1B")}
    for c, name in enumerate(COLS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = center; cell.border = border
    for r, row in enumerate(results, start=2):
        lvl = row.get("_level")
        for c, name in enumerate(COLS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(name, ""))
            cell.border = border
            if lvl in level_fill:
                cell.fill = level_fill[lvl]; cell.font = level_font[lvl]
    for c, w in enumerate([32, 14, 11, 14, 13, 42, 28, 14, 34], start=1):
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:" + chr(64 + len(COLS)) + str(len(results) + 1)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# ---------- routes ----------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/check", methods=["POST"])
def check():
    f = request.files.get("file")
    pasted = (request.form.get("sites") or "").strip()
    sites = []
    if f and f.filename:
        try:
            sites = read_sites_from_bytes(f.filename, f.read())
        except Exception as e:
            return jsonify({"error": "Could not read file: " + str(e)}), 400
    elif pasted:
        seen = set()
        for line in pasted.splitlines():
            h = cert_core.normalize_host(line)
            if h and h not in seen:
                seen.add(h); sites.append(h)
    if not sites:
        return jsonify({"error": "No websites found. Upload a file or paste one site per line."}), 400

    # check concurrently (bounded) for speed
    results = [None] * len(sites)
    def work(idx_host):
        i, host = idx_host
        results[i] = cert_core.classify_site(host)
    with ThreadPoolExecutor(max_workers=min(8, len(sites))) as ex:
        list(ex.map(work, list(enumerate(sites))))

    summary = {"total": len(results), "green": 0, "yellow": 0, "red": 0}
    for r in results:
        summary[r.get("_level", "red")] = summary.get(r.get("_level", "red"), 0) + 1
    return jsonify({"results": results, "summary": summary})


@app.route("/report", methods=["POST"])
def report():
    data = request.get_json(force=True)
    results = data.get("results", [])
    if not results:
        return jsonify({"error": "No results to export."}), 400
    if not HAVE_OPENPYXL:
        # CSV fallback
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=COLS, extrasaction="ignore")
        w.writeheader()
        for r in results: w.writerow(r)
        mem = io.BytesIO(buf.getvalue().encode("utf-8")); mem.seek(0)
        return send_file(mem, as_attachment=True, download_name="certificate_report.csv",
                         mimetype="text/csv")
    buf = build_xlsx(results)
    return send_file(buf, as_attachment=True, download_name="certificate_report.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/health")
def health():
    return {"status": "ok"}


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port, debug=False)
